from pathlib import Path
from datetime import datetime, timedelta
import csv
import json
import shutil
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / 'client/public'
OUT = Path('/home/ubuntu/research')
DATA = json.loads((ROOT / 'client/src/data/prefectures.json').read_text())
RULES = json.loads((ROOT / 'client/src/data/rules.json').read_text())
CHECKED = '2026年10月24日 17時台（JST）'


def sec(value):
    if not value:
        return None
    h, m, s = map(int, value.split(':'))
    return h * 3600 + m * 60 + s


def jp(value):
    if not value:
        return '記録未確認'
    h, m, s = map(int, value.split(':'))
    return f'{h}時間{m:02d}分{s:02d}秒'


def delta(value):
    return f'+{value // 60}分{value % 60:02d}秒'


def source_for(pref, team):
    if team.get('source_url'):
        return team['source_url'], team.get('evidence_level', '補助資料')
    if pref['id'] == 'yamagata':
        return pref['sources'][0]['url'], '公式結果PDF'
    if pref['id'] == 'iwate':
        return 'https://blog.neet-shikakugets.com/iwate-koukou-ekiden-results', '補助結果・2件照合'
    if pref['id'] == 'akita' and team['rank'] >= 2:
        return 'http://www.akiko-rco.org/topics/topics.cgi', '学校OB会資料'
    if pref['id'] == 'miyagi' and team['rank'] >= 2:
        return 'https://www.sposoku.com/koukouekiden-miyagi/', '補助結果資料'
    news = next((s for s in pref['sources'] if 'news' in s['type'] or '報道' in s['type']), pref['sources'][0])
    return news['url'], '報道資料'


rows, missing, pending, extras = [], [], [], []
for pref in DATA:
    first = next((t for t in pref['teams'] if t['rank'] == 1), None)
    if not pref['teams']:
        pending.append({'prefecture': pref['prefecture'], 'event_date': pref['event_date'], 'note': '未開催。2〜4位の3枠は未確定'})
    for team in pref['teams']:
        if team['rank'] not in (2, 3, 4):
            continue
        url, level = source_for(pref, team)
        item = {
            'prefecture_id': pref['id'], 'prefecture': pref['prefecture'],
            'school': team['school'], 'prefecture_rank': team['rank'],
            'time': team['time'], 'time_jp': jp(team['time']), 'seconds': sec(team['time']),
            'source_url': url, 'evidence_level': level, 'note': team['note'],
            'first_gap_seconds': sec(team['time']) - sec(first['time']) if team['time'] and first and first['time'] else None,
        }
        if team['time'] and not team['is_combined']:
            rows.append(item)
        else:
            item['note'] += ' 合同・OP区分確認待ち。' if team['is_combined'] else ''
            missing.append(item)
    # 開催県追加OP、参加数による追加枠は比較本体から区別する。
    if pref['id'] in ('yamagata', 'miyagi'):
        fifth = next((t for t in pref['teams'] if t['rank'] == 5), None)
        if fifth:
            url, level = source_for(pref, fifth)
            extras.append({
                'prefecture': pref['prefecture'], 'school': fifth['school'], 'prefecture_rank': 5,
                'time': fifth['time'], 'time_jp': jp(fifth['time']), 'source_url': url,
                'note': '開催県OP枠の実際の出場校は未確定。この5位校が出場するとは断定しない。' if pref['id'] == 'yamagata' else '県予選15チーム以上の追加正式枠に関する参考。実際の東北大会出場は最終名簿確認待ち。'
            })
rows.sort(key=lambda t: (t['seconds'], t['prefecture'], t['prefecture_rank']))
for index, row in enumerate(rows, 1):
    row['reference_rank'] = index
    row['leader_gap_seconds'] = row['seconds'] - rows[0]['seconds']

refs = []
def ref(url, title):
    for index, item in enumerate(refs, 1):
        if item['url'] == url:
            return index
    refs.append({'url': url, 'title': title})
    return len(refs)

rule_ref = ref(RULES['sources'][0]['url'], '2026年東北高校駅伝 公式実施要項')
national_ref = ref(RULES['sources'][2]['url'], '2026年全国高校駅伝 公式大会要項')
for row in rows + missing + extras:
    row['reference_id'] = ref(row['source_url'], f"{row['prefecture']} 2026年女子県高校駅伝の結果資料")
for pref in DATA:
    if not pref['teams']:
        ref(pref['sources'][0]['url'], f"{pref['prefecture']} 2026年女子県高校駅伝の公式日程・要項")
ref('https://www.sposoku.com/koukouekiden-iwate/', '岩手県2026年女子結果の照合用資料')
ref('https://blog.neet-shikakugets.com/miyagi-koukou-ekiden-results', '宮城県2026年女子2位の照合用資料')

payload = {'checked_at': CHECKED, 'ranked': rows, 'unranked': missing, 'pending': pending, 'additional_reference': extras, 'references': refs}
(ROOT / 'client/src/data/comparison.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')
(PUBLIC / 'comparison.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')

summary = []
if len(rows) >= 2:
    a, b = rows[:2]
    summary.append(f"確認済みの参考タイムでは、{a['school']} {a['time_jp']}、{b['school']} {b['time_jp']}が先頭で、差は{b['leader_gap_seconds']}秒です。これは県大会の記録順で、東北大会の予想順位ではありません。")
third_fourth = [r for r in rows if r['prefecture_rank'] in (3, 4)]
if third_fourth:
    best = third_fourth[0]
    summary.append(f"県3・4位校の中で最も速い確認記録は{best['prefecture']}の{best['school']}で、{best['time_jp']}です。全体参考{best['reference_rank']}番目に入るため、県2位だけの比較では対象校の見え方が変わります。")
akita = {r['prefecture_rank']: r for r in rows if r['prefecture_id'] == 'akita'}
if 2 in akita and 3 in akita:
    diff = akita[3]['seconds'] - akita[2]['seconds']
    summary.append(f"秋田北鷹と秋田工業の差は{diff}秒です。県3位の秋田工業も含め、当日のオーダーや区間ごとの走りで逆転し得る候補として見るのが適切です。")
miyagi4 = next((r for r in rows if r['prefecture_id'] == 'miyagi' and r['prefecture_rank'] == 4), None)
if miyagi4 and 2 in akita:
    close = abs(miyagi4['seconds'] - akita[2]['seconds'])
    summary.append(f"宮城4位の古川学園と秋田2位の秋田北鷹の記録差は{close}秒です。県内順位が違っても総合記録は近く、県4位校を除くと比較から漏れる例です。")
summary.append('山形県予選と東北大会は同じ長井市の会場名・周回コースで案内されていますが、区間設定や当日の条件を同一と決めつけることはできません。県間の総合タイムから勝率を算出したり、出場権の獲得を断定したりはしていません。')
summary.append('青森・福島の予選が終わるまでは全県の比較が完成していません。比較対象の県2〜4位校も、正式な東北大会出場名簿と実際の県代表校を確認してから全国地区代表の候補を確定します。')

payload['analysis'] = summary
(ROOT / 'client/src/data/comparison.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')
(PUBLIC / 'comparison.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')

lines = ['# 2026年女子高校駅伝｜県2〜4位の参考タイム順', '', f'**確認日：{CHECKED}**。各県の県大会1位を除き、2〜4位の確認できた総合タイムを速い順に並べています。これは東北大会の予想順位・正式出場校一覧ではありません。', '', '|参考順|県|高校名|県内順位|総合タイム|先頭との差|確認資料|', '|---:|---|---|---:|---|---:|---|']
for row in rows:
    lines.append(f"|{row['reference_rank']}|{row['prefecture']}|{row['school']}|{row['prefecture_rank']}位|{row['time_jp']}|{'—' if row['leader_gap_seconds'] == 0 else delta(row['leader_gap_seconds'])}|{row['evidence_level']} [{row['reference_id']}]|")
if missing:
    lines += ['', '## タイム未確認の比較対象校', '', '|県|高校名|県内順位|状況|確認資料|', '|---|---|---:|---|---|']
    for row in missing:
        lines.append(f"|{row['prefecture']}|{row['school']}|{row['prefecture_rank']}位|学校名・順位は確認、記録は未確認。参考順には入れない。|[{row['reference_id']}]|")
lines += ['', '青森は10月29日、福島は10月31日の開催予定です。両県の2〜4位は未確定のため、学校名・記録を代用していません。']
lines += ['', '## 正式出場枠と山形の開催県追加枠', '', f'公式要項10参加資格⑵では女子は各県上位4チームで、県予選参加数が15チーム以上の県は1チーム追加できます。開催県山形については10⑻に「オープン参加を各1チーム認める」とあります。開催県の追加1チームと、参加数による正式枠の追加は別の規定です。 [{rule_ref}]', '', f'全国の地区代表は東北大会において実際の県代表校を除く最上位校1校です。今回はご指定どおり県2〜4位を比較の本体とし、開催県OPの実際の出場校や全国選考資格を未確認のまま確定扱いしていません。OPの扱いは主催者の正式発表・監督会議資料を確認する必要があります。 [{rule_ref}] [{national_ref}]']
if extras:
    lines += ['', '|追加枠に関する参考|学校|県内順位|タイム|扱い|', '|---|---|---:|---|---|']
    for row in extras:
        lines.append(f"|{row['prefecture']}|{row['school']}|5位|{row['time_jp']}|{row['note']} [{row['reference_id']}]|")
lines += ['', '## 2位だけでなく3・4位も含めて読む', '']
for text in summary:
    used = sorted({r['reference_id'] for r in rows})
    citations = ' '.join(f'[{i}]' for i in used) if ('秒' in text or '確認記録' in text) else f'[{rule_ref}] [{next(r["reference_id"] for r in rows if r["prefecture_id"] == "yamagata")}]' if '会場' in text else ''
    lines += [text + (' ' + citations if citations else ''), '']
lines += ['## 県ごとの1〜4位', '', '|県|1位・比較対象外|2位|3位|4位|', '|---|---|---|---|---|']
for pref in DATA:
    cells = []
    for rank in range(1, 5):
        team = next((t for t in pref['teams'] if t['rank'] == rank), None)
        cells.append(f"{team['school']} {team['time'] or '記録未確認'}" if team else '未開催・未確定' if pref['status'] == 'scheduled' else '未確認')
    source = pref['sources'][0]['url'] if not pref['teams'] else source_for(pref, next(t for t in pref['teams'] if t['rank'] == 2))[0]
    source_id = next(i for i, item in enumerate(refs, 1) if item['url'] == source)
    lines.append(f"|{pref['prefecture']} [{source_id}]|" + '|'.join(cells) + '|')
lines += ['', '記録の根拠は公式PDF、報道、学校関係資料、結果速報を区別して表記しています。補助資料による確認は公式全成績での確認とは異なります。県代表の辞退等がある場合、東北大会で除外するのは実際に県代表となった校です。', '', '調査・制作：Manus AI', '', '## References', '']
for index, item in enumerate(refs, 1):
    url = item['url'].replace('(', '%28').replace(')', '%29')
    lines.append(f'[{index}]: {url} "{item["title"]}"')
report = '\n'.join(lines) + '\n'
(OUT / 'contenders-2026.md').write_text(report)
(PUBLIC / 'report.md').write_text(report)

text = ['2026年 女子高校駅伝｜県2〜4位 参考タイム順', f'確認日：{CHECKED}', '県1位は除外。県間の参考記録順であり、東北大会の予想順位ではありません。', '']
for row in rows:
    text.append(f"{row['reference_rank']:>2}位　{row['prefecture']}　{row['school']}　{row['time_jp']}　（県{row['prefecture_rank']}位）")
for row in missing:
    text.append(f"参考順なし　{row['prefecture']}　{row['school']}　記録未確認　（県{row['prefecture_rank']}位）")
text += ['', '未開催：青森県10/29、福島県10/31。両県2〜4位は未確定。', '', '追加枠：各県通常4チーム、県予選15チーム以上なら1チーム追加。山形の開催県追加1チームはOP参加枠。実際のOP出場校と全国選考資格は確定扱いしない。', '']
for row in extras:
    text.append(f"追加枠参考：{row['prefecture']} {row['school']} {row['time_jp']}（県5位）。{row['note']}")
text += ['', '考察'] + summary + ['', '出典']
for index, item in enumerate(refs, 1):
    text += [f'[{index}] {item["title"]}', item['url']]
plain = '\n'.join(text) + '\n'
(OUT / 'contenders-2026.txt').write_text(plain)
(PUBLIC / 'comparison.txt').write_text(plain)

with (PUBLIC / 'comparison.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(['参考順', '県', '高校名', '県内順位', '総合タイム', '日本語タイム', '先頭との差', '確認資料', '出典URL'])
    for row in rows:
        writer.writerow([row['reference_rank'], row['prefecture'], row['school'], row['prefecture_rank'], row['time'], row['time_jp'], delta(row['leader_gap_seconds']), row['evidence_level'], row['source_url']])
    for row in missing:
        writer.writerow(['', row['prefecture'], row['school'], row['prefecture_rank'], '', '記録未確認', '', row['evidence_level'], row['source_url']])

wb = Workbook()
wb.remove(wb.active)
NAVY, GREEN, RUST, LIGHT = '253D31', 'E8F1EA', 'C63E2C', 'F5F7F4'
DURATION = '[h]"時間"mm"分"ss"秒"'
DIFF = '+m"分"ss"秒"'


def sheet(name, title, subtitle, headers, widths):
    ws = wb.create_sheet(name)
    n = len(headers)
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws['A1'] = title
    ws['A1'].font = Font(name='游ゴシック', size=20, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 42
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws['A2'] = subtitle
    ws['A2'].font = Font(name='游ゴシック', size=10, color='556459')
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 35
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n)
    ws['A3'] = '確認日：' + CHECKED + '　／　公式結果と補助資料を区別。県間の記録順は予測順位ではありません。'
    ws['A3'].font = Font(name='游ゴシック', size=9, color='69786C')
    ws['A3'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[3].height = 28
    ws.append([])
    for c, value in enumerate(headers, 1):
        cell = ws.cell(5, c, value)
        cell.font = Font(name='游ゴシック', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[5].height = 29
    for c, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = 'D6'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:5'
    return ws


def finish(ws, rows_data, table_name, time_cols=(), links=()):
    for r, values in enumerate(rows_data, 6):
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name='游ゴシック', size=11, color=NAVY)
            cell.fill = PatternFill('solid', fgColor=LIGHT if r % 2 == 0 else 'FFFFFF')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = Border(bottom=Side(style='hair', color='D7DFD5'))
            if c in time_cols and value is not None and not isinstance(value, str):
                cell.number_format = DURATION
            if c in links and value:
                cell.hyperlink = value
                cell.value = '出典を開く'
                cell.font = Font(name='游ゴシック', size=10, color='26734D', underline='single')
        ws.row_dimensions[r].height = 32
    if rows_data:
        table = Table(displayName=table_name, ref=f'A5:{get_column_letter(len(rows_data[0]))}{5 + len(rows_data)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showFirstColumn=False, showLastColumn=False, showRowStripes=True)
        ws.add_table(table)
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'


ws = sheet('参考タイム順', '県2〜4位｜参考タイム順', '県1位は除外。記録を確認できた学校のみを速い順に表示。正式出場校は最終エントリー確認待ち。', ['参考順', '県', '高校名', '県内順位', '総合タイム', '先頭との差', '県1位との差', '確認資料', '出典'], [10, 12, 34, 12, 21, 16, 18, 24, 16])
values = []
for r, row in enumerate(rows, 6):
    values.append([row['reference_rank'], row['prefecture'], row['school'], row['prefecture_rank'], row['seconds'] / 86400, row['leader_gap_seconds'] / 86400, row['first_gap_seconds'] / 86400 if row['first_gap_seconds'] is not None else None, row['evidence_level'], row['source_url']])
finish(ws, values, 'ContenderTimes', time_cols=(5, 7), links=(9,))
for r, row in enumerate(rows, 6):
    ws.cell(r, 6).number_format = DIFF
    ws.cell(r, 7).number_format = DIFF
    if row['prefecture_rank'] in (3, 4):
        ws.cell(r, 4).fill = PatternFill('solid', fgColor=GREEN)
        ws.cell(r, 4).font = Font(name='游ゴシック', bold=True, color='276C49')
    if r <= 8:
        ws.cell(r, 1).font = Font(name='游ゴシック', size=13, bold=True, color=RUST)

ws = sheet('県別1〜4位', '6県の県大会1〜4位', '県大会1位は比較対象外。未開催・タイム未確認は空欄をゼロにせず、その状況を表示。', ['県', '県内順位', '高校名', '総合タイム', '比較対象', '確認資料', '出典', '備考'], [12, 12, 34, 22, 16, 24, 16, 60])
values = []
for pref in DATA:
    for rank in range(1, 5):
        team = next((t for t in pref['teams'] if t['rank'] == rank), None)
        if team:
            url, level = source_for(pref, team)
            values.append([pref['prefecture'], rank, team['school'], sec(team['time']) / 86400 if team['time'] else '記録未確認', '対象外・県1位' if rank == 1 else '2〜4位対象', level, url, team['note']])
        else:
            values.append([pref['prefecture'], rank, '未開催・未確定' if pref['status'] == 'scheduled' else '未確認', None, '結果確定待ち', '公式日程・要項', pref['sources'][0]['url'], pref['event_date'] + '開催予定' if pref['status'] == 'scheduled' else '今回未確認'])
finish(ws, values, 'PrefecturalTopFour', time_cols=(4,), links=(7,))
for r, value_row in enumerate(values, 6):
    ws.row_dimensions[r].height = max(32, min(112, ((len(value_row[7]) + 27) // 28) * 14 + 8))
for r, values_row in enumerate(values, 6):
    if values_row[1] == 1:
        for cell in ws[r]:
            cell.font = Font(name='游ゴシック', size=10, color='7F897F')

ws = sheet('追加枠・未確認', '追加枠と未確認情報を区別', '山形の開催県追加1チームはOP枠。県予選15チーム以上の正式枠追加とは別。実際の出場校・OPの全国選考資格は主催者確認前に確定扱いしない。', ['区分', '県', '高校名', '県内順位', '総合タイム', '扱い・備考', '出典'], [22, 12, 34, 12, 22, 74, 16])
values = [['参加資格', '東北6県', '各県通常4チーム', None, None, '県予選15チーム以上の県は1チーム追加できる。', RULES['sources'][0]['url']], ['開催県OP', '山形県', '実際の出場校は未確定', None, None, '開催県から女子1チームのOP参加を認める。通常上位4枠とは別。', RULES['sources'][0]['url']]]
for row in extras:
    values.append(['追加枠参考', row['prefecture'], row['school'], 5, sec(row['time']) / 86400 if row['time'] else '記録未確認', row['note'], row['source_url']])
for row in missing:
    values.append(['タイム未確認', row['prefecture'], row['school'], row['prefecture_rank'], '記録未確認', '参考順位には入れない。' + row['note'], row['source_url']])
for row in pending:
    values.append(['未開催', row['prefecture'], '2〜4位は未確定', None, None, row['event_date'] + '開催予定。過年度の記録は代用しない。', next(p['sources'][0]['url'] for p in DATA if p['prefecture'] == row['prefecture'])])
finish(ws, values, 'ExceptionsAndPending', time_cols=(5,), links=(7,))
for r in range(6, ws.max_row + 1):
    ws.row_dimensions[r].height = max(48, min(120, ((len(str(ws.cell(r, 6).value)) + 34) // 35) * 14 + 8))

ws = sheet('出典と考察', '資料リンクと比較の読み方', '結果本文の確認レベルを区別し、県2位だけでなく3・4位も含めて比較。', ['番号・区分', '内容', 'URL'], [18, 104, 22])
values = [[f'資料 {i}', item['title'], item['url']] for i, item in enumerate(refs, 1)]
values += [['考察', t, None] for t in summary]
finish(ws, values, 'SourcesAndNotes', links=(3,))
for r in range(6, ws.max_row + 1):
    ws.row_dimensions[r].height = 54 if r > 5 + len(refs) else 34

xlsx = OUT / 'contenders-2026.xlsx'
wb.save(xlsx)
shutil.copy2(xlsx, PUBLIC / 'comparison.xlsx')
shutil.copy2(PUBLIC / 'comparison.csv', OUT / 'contenders-2026.csv')
check = load_workbook(xlsx)
assert len(check.worksheets) == 4
assert check['参考タイム順'].max_row == 5 + len(rows)
assert check['県別1〜4位'].max_row == 29
print(f'比較表を生成：記録確認 {len(rows)}校、記録未確認 {len(missing)}校、未開催 {len(pending)}県。')
print(f'Excel: {xlsx}')
